"""Nájemní smlouva a dodatek ve Word formátu, ve stejné struktuře jako vzorová smlouva k bytu Revoluční."""
import json
from datetime import date
from pathlib import Path

from docx.enum.text import WD_BREAK
from docx.shared import Pt

from . import docgen

TEXT = json.loads((Path(__file__).parent / "templates" / "najem_text.json").read_text(encoding="utf-8"))
PAY_DEFAULT = "28. den v měsíci na měsíc následující (nájemné předem)"
DEPOSIT_DEFAULT = "Jistota byla uhrazena při podpisu okamžitým převodem na účet pronajímatele na účet dle #8."


def cz_date(v) -> str:
    if not v:
        return "—"
    d = v if isinstance(v, date) else date.fromisoformat(str(v)[:10])
    return f"{d.day}. {d.month}. {d.year}"


def money(v) -> str:
    return f"{round(float(v or 0)):,}".replace(",", ".") + ",- Kč"


ROMAN = ["", "I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]


def _adv_total(adv: dict) -> float:
    return sum(float(v or 0) for v in adv.values())


def _party(d: docgen.DocxDocument, label: str, p: dict, lines: list[str]) -> None:
    d.add_paragraph().add_run(label).bold = True
    for ln in lines:
        if ln:
            d.add_paragraph(ln)


def _clauses(d, paras: list[str]) -> None:
    for line in paras:
        if line.startswith("Článek "):
            p = d.add_paragraph()
            r = p.add_run(line)
            r.bold = True
        else:
            d.add_paragraph(line)


def lease_docx(flat, b: dict) -> bytes:
    """Nájemní smlouva. b = údaje z formuláře, flat = byt, adv = zálohy v Kč."""
    adv = {k: float(v) for k, v in (b.get("advances") or {}).items() if float(v or 0)}
    rent = float(b.get("rent") or 0)
    deposit = float(b.get("deposit") or 0)
    d = docgen._new("Nájemní smlouva")
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_paragraph("dle § 2235 a násl. zákona č. 89/2012 Sb., občanského zákoníku, ve znění pozdějších předpisů (dále jen občanský zákoník)")
    body = TEXT["uvod_a_clanky"]
    first_art = body.index("Článek I.")
    d.add_heading("Úvod", level=2)
    for line in body[1:first_art]:
        d.add_paragraph(line)

    L, T = b.get("landlord") or {}, b.get("tenant") or {}
    d.add_heading("1. SMLUVNÍ STRANY", level=2)
    t = d.add_table(rows=1, cols=2)
    t.style = "Table Grid"
    t.rows[0].cells[0].text = "Pronajímatel:\n" + "\n".join(x for x in [
        L.get("name"), L.get("address") and f"Trvale bytem: {L.get('address')}",
        (L.get("phone") or L.get("email")) and f"Telefon: {L.get('phone','')}, e-mail: {L.get('email','')}"] if x)
    t.rows[0].cells[1].text = "Nájemce:\n" + "\n".join(x for x in [
        T.get("name"), T.get("birth") and f"Narozen(a): {T.get('birth')}",
        T.get("address") and f"Trvale bytem: {T.get('address')}",
        T.get("op") and f"Číslo OP: {T.get('op')}",
        (T.get("phone") or T.get("email")) and f"Telefon: {T.get('phone','')}, e-mail: {T.get('email','')}"] if x)
    d.add_paragraph()

    d.add_heading("2. PARAMETRY SMLOUVY", level=2)
    doba = "určitou" if b.get("valid_to") and not b.get("indefinite") else "neurčitou"
    rows = [
        ("Adresa nemovitosti", f"{flat.addr}, {flat.zip or ''}", "#1"),
        ("Číslo bytové jednotky", b.get("unit") or "", "#1"),
        ("Podlaží a umístění", b.get("floor") or flat.floor or "", "#1"),
        ("Výměra a dispozice", f"{b.get('area') or ''} m2, {b.get('dispo') or flat.dispo or ''}", "#1"),
        ("Datum uzavření smlouvy", cz_date(b.get("contract_date")), "#2"),
        ("Doba smlouvy", doba, "#3"),
        ("Počátek nájmu", cz_date(b.get("valid_from")), "#4"),
        ("Ukončení nájmu (u pronájmu na dobu určitou)", "" if doba == "neurčitou" else cz_date(b.get("valid_to")), "#5"),
        ("Měsíční nájem vč. záloh", money(rent + _adv_total(adv)), "#6"),
        ("Variabilní symbol", b.get("vs") or "", "#7"),
        ("Číslo účtu pronajímatele", b.get("bank_acc") or flat.bank_acc or "", "#8"),
        ("Splatnost nájmu", b.get("payment_due") or PAY_DEFAULT, "#9"),
        ("Výše jistoty", money(deposit), "#10"),
        ("Způsob úhrady jistoty", b.get("deposit_note") or DEPOSIT_DEFAULT, "#11"),
        ("Další ujednání", b.get("other") or "", "#12"),
    ]
    docgen._table(d, ["Položka", "Údaj", "Odkaz"], [list(r) for r in rows])

    d.add_heading("3. VÝPOČET ZÁLOH NA SLUŽBY", level=2)
    srows = [[k, money(v)] for k, v in adv.items()]
    srows.append(["Celkem služby", money(_adv_total(adv))])
    srows.append(["Čisté nájemné", money(rent)])
    docgen._table(d, ["Služba", "Měsíčně"], srows, money_cols=(1,))

    d.add_heading("4. OSOBY NAHLÁŠENÉ K POBYTU V NEMOVITOSTI", level=2)
    persons = [p for p in (b.get("persons") or []) if str(p).strip()]
    d.add_paragraph(f"Celkový počet osob: {len(persons) or 1}")
    for p in persons or [T.get("name") or ""]:
        d.add_paragraph(str(p))
    d.add_paragraph()

    _clauses(d, body[first_art:])
    d.add_paragraph()
    d.add_paragraph("______________________          ______________________")
    d.add_paragraph("Pronajímatel                     Nájemce")
    d.add_page_break()
    for line in TEXT["protokol"]:
        p = d.add_paragraph()
        r = p.add_run(line)
        if line.isupper() or line.startswith("VZNIK") or line.startswith("UKONČENÍ"):
            r.bold = True
    return docgen._save(d)


def amendment_docx(flat, b: dict, old: dict) -> bytes:
    """Dodatek. old = předchozí předpis (rent, deposit, advances, valid_from, valid_to)."""
    eff = cz_date(b.get("eff_date"))
    ch = b.get("changes") or {}
    d = docgen._new(f"Dodatek č. {b.get('amendment_no') or '[doplňte]'} ke smlouvě o nájmu bytu")
    d.add_paragraph(
        f"uzavřený dle § 2235 a násl. zákona č. 89/2012 Sb., občanského zákoníku, mezi smluvními stranami:")
    L, T = b.get("landlord") or {}, b.get("tenant") or {}
    d.add_paragraph(f"Pronajímatel: {L.get('name') or '[doplňte]'}, trvale bytem {L.get('address') or '[doplňte]'}")
    d.add_paragraph(f"Nájemce: {T.get('name') or '[doplňte]'}, trvale bytem {T.get('address') or '[doplňte]'}")
    d.add_paragraph(
        f"Smluvní strany uzavřely dne {cz_date(b.get('contract_date'))} nájemní smlouvu k bytu "
        f"{flat.addr}, {flat.zip or ''} (dále jen „smlouva“).")
    n = 0
    arts = []

    def art(title, text):
        nonlocal n
        n += 1
        arts.append(f"Článek {ROMAN[n]}.\n{title}")
        arts.append(text)

    if ch.get("rent"):
        new_rent = float(b.get("rent") or 0)
        art("Výše nájmu", f"Nájemné bez záloh se od {eff} mění z {money(old.get('rent'))} na {money(new_rent)} měsíčně.")
    if ch.get("adv"):
        new_adv = {k: float(v) for k, v in (b.get("advances") or {}).items() if float(v or 0)}
        lst = "; ".join(f"{k}: {money(v)}" for k, v in new_adv.items()) or "bez záloh"
        art("Zálohy na služby",
            f"Měsíční zálohy na služby se od {eff} mění takto: {lst}. Celkem {money(_adv_total(new_adv))} měsíčně. "
            "Článek II smlouvy a oddíl 3 smlouvy se tím nahrazují.")
    if ch.get("term"):
        to = b.get("valid_to") and not b.get("indefinite")
        art("Doba nájmu",
            f"Nájem se mění tak, že trvá od {cz_date(b.get('valid_from'))} do "
            f"{cz_date(b.get('valid_to')) if to else 'neurčité doby'}.")
    if ch.get("deposit"):
        art("Jistota", f"Výše jistoty se mění z {money(old.get('deposit'))} na {money(b.get('deposit'))}.")
    if ch.get("other") and (b.get("other") or "").strip():
        art("Jiná ujednání", b["other"].strip())
    if not arts:
        art("Změny", "Žádné změny.")
    for line in arts:
        p = d.add_paragraph()
        r = p.add_run(line)
        if line.startswith("Článek "):
            r.bold = True
    p = d.add_paragraph()
    p.add_run(f"Článek {ROMAN[n + 1]}.").bold = True
    d.add_paragraph("Závěrečná ustanovení")
    d.add_paragraph(
        "Ostatní ustanovení smlouvy zůstávají beze změny. Tento dodatek nabývá účinnosti dnem "
        f"{eff}. Je vyhotoven ve dvou stejnopisech s platností originálu, po jednom pro každou smluvní stranu.")
    d.add_paragraph()
    d.add_paragraph("______________________          ______________________")
    d.add_paragraph("Pronajímatel                     Nájemce")
    return docgen._save(d)
