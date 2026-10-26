"""Předávací protokol bytu ve Word formátu: stavy měřidel, volitelně s fotkami jako přílohami."""
from io import BytesIO

from docx.shared import Mm

from . import docgen
from .contracts import _signatures, cz_date


def handover_docx(flat, b: dict) -> bytes:
    d = docgen._new("Předávací protokol")
    d.add_paragraph(f"Předávací protokol k bytu {flat.addr}, {flat.zip or ''}")
    d.add_paragraph(f"Datum předání: {cz_date(b.get('date'))}")
    d.add_paragraph(f"Pronajímatel: {b.get('landlord_name') or ''}")
    d.add_paragraph("Nájemce: " + (", ".join(p for p in (b.get("people") or []) if p) or "—"))
    d.add_heading("Stavy měřidel", level=2)
    rows = [[m.get("type", ""), m.get("room", "") or "—", m.get("value") or "nezapsáno", m.get("photo") or "bez fotky"]
            for m in (b.get("meters") or [])]
    docgen._table(d, ["Měřidlo", "Místnost", "Stav", "Fotka"], rows or [["—", "—", "—", "—"]])
    d.add_heading("Poznámky ke stavu bytu a předání klíčů", level=2)
    d.add_paragraph((b.get("notes") or "").strip() or "Bez zvláštních poznámek.")
    d.add_paragraph()
    d.add_paragraph("Nájemce potvrzuje převzetí bytu ve stavu popsaném výše a stavy měřidel k datu předání.")
    _signatures(d, b)
    photos = b.get("_photos") or []
    if photos:
        d.add_page_break()
        d.add_heading("Přílohy: fotky měřidel", level=2)
        for name, data in photos:
            try:
                d.add_paragraph(name)
                d.add_picture(BytesIO(data), width=Mm(130))
            except Exception:
                d.add_paragraph("(fotku se nepodařilo vložit)")
    return docgen._save(d)
